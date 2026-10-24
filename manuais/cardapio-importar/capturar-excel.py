"""Captura a importação por planilha (Cardápio → Produtos → Importar) para o #130.

Baixa o **template de verdade** pelo próprio botão da tela (`expect_download`) e preenche uma
cópia dele com cinco porções. Fazer assim, em vez de montar uma planilha do zero com nomes de
coluna chutados, prova de passagem que o mapeamento automático reconhece os cabeçalhos que o
sistema distribui — se um dia o template mudar, este script acusa.

O cadastro no fim é real, e foi autorizado pelo dono em 23/10/2026.

Requer `openpyxl` (`pip install openpyxl`).
"""
import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "painel-entregador"))

import openpyxl
from playwright.sync_api import sync_playwright
from beefood import abrir, after_click, limpar_tela

AQUI = Path(__file__).resolve().parent
PURAS = AQUI / "imagens-puras"
TEMPLATE = Path("/tmp/template_importacao_produtos.xlsx")
PLANILHA = Path("/tmp/porcoes-para-importar.xlsx")

# Setor, Descrição, Código, Preço Venda, Preço Custo, Descritivo — na ordem do template.
PORCOES = [
    ("Porções", "Batata Frita Pequena", "POR001", 18.00, 6.50, "300g de batata rústica com páprica"),
    ("Porções", "Batata Frita Grande", "POR002", 29.00, 11.00, "600g de batata rústica com páprica"),
    ("Porções", "Onion Rings", "POR003", 24.00, 9.00, "12 anéis de cebola empanados"),
    ("Porções", "Nuggets 10 Unidades", "POR004", 21.00, 8.00, "10 nuggets de frango com molho"),
    ("Porções", "Mix de Porções", "POR005", 39.00, 15.00, "batata, onion rings e nuggets"),
]


def preencher_planilha():
    shutil.copy(TEMPLATE, PLANILHA)
    wb = openpyxl.load_workbook(PLANILHA)
    ws = wb.active
    cabecalhos = [c.value for c in ws[1]]
    print("cabeçalhos do template:", cabecalhos)
    indice = {nome: i + 1 for i, nome in enumerate(cabecalhos)}
    for linha, (setor, nome, codigo, venda, custo, descritivo) in enumerate(PORCOES, start=2):
        ws.cell(linha, indice["Setor"], setor)
        ws.cell(linha, indice["Descrição"], nome)
        ws.cell(linha, indice["Código"], codigo)
        ws.cell(linha, indice["Preço Venda"], venda)
        ws.cell(linha, indice["Preço Custo"], custo)
        ws.cell(linha, indice["Descritivo"], descritivo)
        ws.cell(linha, indice["Unidade Medida"], "UN")
    wb.save(PLANILHA)
    print("planilha preenchida:", PLANILHA, f"{len(PORCOES)} linhas")


with sync_playwright() as p:
    browser, ctx, page = abrir(p, "/cardapio?tab=produtos", viewport=(1440, 900), escala=1.5)
    ctx.set_default_timeout(30000)
    after_click(page, 8000)
    limpar_tela(page)

    # 13. O passo 1 do modal: onde o arquivo entra e de onde sai o template.
    page.locator("button").filter(has_text="Importar").first.click()
    after_click(page, 3000)
    dialogo = page.locator('div[role="dialog"]').filter(has_text="Selecione o arquivo").last
    page.screenshot(path=str(PURAS / "13-excel-passo1.png"))
    print("13 passo 1")

    with page.expect_download() as baixado:
        dialogo.locator("button").filter(has_text="BAIXAR TEMPLATE").first.click()
    baixado.value.save_as(str(TEMPLATE))
    print("template baixado:", TEMPLATE, TEMPLATE.stat().st_size, "bytes")
    preencher_planilha()

    dialogo.locator('input[type="file"]').first.set_input_files(str(PLANILHA))
    after_click(page, 4000)

    # 14. O mapeamento: a tela já associou cada coluna sozinha, e mostra um exemplo de cada.
    dialogo = page.locator('div[role="dialog"]').filter(has_text="Mapeamento de colunas").last
    page.screenshot(path=str(PURAS / "14-excel-mapeamento.png"))
    print("14 mapeamento")
    print(dialogo.inner_text()[:700])

    dialogo.locator("button").filter(has_text="IMPORTAR").first.click()
    page.wait_for_timeout(2500)
    page.screenshot(path=str(PURAS / "15-excel-importando.png"))
    print("15 importando")

    for i in range(40):
        page.wait_for_timeout(2000)
        texto = page.locator('div[role="dialog"]').last.inner_text()
        if "Resultado" in texto:
            print(f"  terminou em ~{(i + 1) * 2}s")
            break
    after_click(page, 4000)

    # 16. O resultado, linha por linha — é aqui que se vê o que não entrou.
    page.screenshot(path=str(PURAS / "16-excel-resultado.png"))
    print("16 resultado")
    print(page.locator('div[role="dialog"]').last.inner_text()[:600])
    browser.close()
print("fim")
